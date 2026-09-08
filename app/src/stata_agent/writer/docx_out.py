"""Word 初稿最小版（切片 5b）：把"结果句（claim+数字已接地）"渲染成 .docx。

溯源策略先做轻量：每句标题/正文尾部附 [claim-xxx] 小注（学者可见、可删）；
"隐藏批注/单独引用清单"留给切片 5 完整版（DD-05 §6）。
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document

from ..domain.models import Claim


def claims_to_docx(
    title: str,
    paragraphs: list[tuple[str, Claim | None]],
    *,
    intro: str = "实证结果初稿（切片 5b 自动渲染，数字均已接地可溯源）。",
) -> BytesIO:
    """paragraphs = [(句子文本, 支撑 claim 或 None)]。返回 .docx 字节流。"""
    doc = Document()
    doc.add_heading(title, level=0)
    doc.add_paragraph(intro)
    for text, claim in paragraphs:
        p = doc.add_paragraph(text)
        if claim is not None:
            note = p.add_run(f"   [{claim.claim_id}]")
            note.font.size = note.font.size  # 保持默认
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def save_docx(buf: BytesIO, path: str | Path) -> Path:
    p = Path(path)
    p.write_bytes(buf.getvalue())
    return p


def tables_to_docx(title: str, tables: list) -> BytesIO:
    """把一组 TableModel 渲染成 .docx（含表头与行）。"""
    from .table import add_table_to_doc

    doc = Document()
    doc.add_heading(title, level=0)
    for model in tables:
        add_table_to_doc(doc, model, title_heading=True)
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
