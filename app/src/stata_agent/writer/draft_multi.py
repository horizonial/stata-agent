"""多表实证初稿（Wave-1 论文化最小）：描述表 + 主表(带星) + 稳健性小结 → .docx。

数据来源：run_variants 的 results（machine: coef/se/N/r2）；数字格可回链 run。
"""

from __future__ import annotations

import csv

from docx import Document

from ..tools.robustness import check_robustness, stars
from .table import Cell, Row, TableModel, add_table_to_doc
from .validation import DeliveryManifestV1


def _f(x, digits=3):
    return f"{x:.{digits}f}" if isinstance(x, (int, float)) else "-"


def regression_table_from_results(results: dict[str, dict]) -> TableModel:
    ok = [(k, v) for k, v in results.items() if v.get("machine")]
    model = TableModel(title="主回归与稳健性", columns=[v.get("label", k) for k, v in ok])
    stats = {
        "系数": lambda m: _f(m.get("coef")) + stars(m.get("coef"), m.get("se")),
        "标准误": lambda m: _f(m.get("se")),
        "样本量": lambda m: str(int(m.get("N", 0))),
        "R²": lambda m: _f(m.get("r2")),
    }
    for label, fmt in stats.items():
        model.rows.append(Row(label=label, cells=[Cell(text=fmt(v["machine"])) for _, v in ok]))
    return model


def table1_from_long(path) -> TableModel | None:
    """ck_long.csv → 描述表：按 (nj,wave) 的 fte 均值（观测数）。"""
    rows = list(csv.DictReader(open(path, encoding="utf-8")))
    vals: dict[tuple[int, int], list[float]] = {}
    for r in rows:
        if r["fte"] not in (None, ""):
            vals.setdefault((int(r["nj"]), int(r["wave"])), []).append(float(r["fte"]))
    group = {k: (sum(v) / len(v), len(v)) for k, v in vals.items()}
    model = TableModel(title="表 1 描述统计：FTE 均值（观测数）", columns=["政策前", "政策后"])
    for nj, name in ((0, "宾州 PA"), (1, "新泽西 NJ")):
        texts = []
        for w in (1, 2):
            g = group.get((nj, w))
            texts.append(f"{g[0]:.2f}（{g[1]}）" if g else "-")
        model.rows.append(Row(label=name, cells=[Cell(text=t) for t in texts]))
    return model


def build_draft_package(
    proj,
    *,
    title: str = "实证研究初稿",
    method: str = "",
    limits: str = "",
    library=None,
    figure_items: list[dict] | None = None,
) -> tuple[bytes, DeliveryManifestV1]:
    """Build a validated DOCX and its derived evidence manifest."""

    from .ground import render_claim_sentence
    from .validation import preflight_delivery
    from .table import build_run_table

    ok_runs = [
        (rid, rec)
        for rid, rec in proj.runs.items()
        if rec.status == "succeeded" and rec.machine
    ]
    run_models = {rid: build_run_table(proj, rid, title="") for rid, _rec in ok_runs}
    active_claims = [
        claim
        for claim in sorted(proj.claims.values(), key=lambda c: c.claim_id)
        if claim.status == "supported"
    ]
    if not ok_runs and not active_claims and not (figure_items or []):
        raise ValueError("没有 active claim 或含 provenance 的 run，不能出稿")

    table_model = None
    table_objects: list[dict] = []
    if ok_runs:
        table_model = TableModel(
            title="回归结果（由事件账本渲染）",
            columns=[rid[:14] for rid, _ in ok_runs],
        )
        labels = {
            "系数": "核心系数",
            "标准误": "SE",
            "样本量": "样本量",
            "R²": "R²",
        }
        for label, source_label in labels.items():
            cells = []
            for rid, _rec in ok_runs:
                source = next(
                    (row.cells[0] for row in run_models[rid].rows if row.label == source_label),
                    None,
                )
                if source is None:
                    cells.append(Cell(text="-", stat_type=label))
                else:
                    cells.append(
                        Cell(
                            text=source.text,
                            stat_type=source.stat_type,
                            numeric=source.numeric,
                            card_id=source.card_id,
                        )
                    )
            table_model.rows.append(Row(label=label, cells=cells))
        from .table import numeric_cells

        for row_index, col_index, cell in numeric_cells(table_model):
            table_objects.append(
                {
                    "object_type": "table_cell",
                    "object_id": f"regression:{row_index}:{col_index}",
                    "card_ids": [cell.card_id] if cell.card_id else [],
                }
            )

    claim_texts: list[dict] = []
    citation_texts: list[dict] = []
    for claim in active_claims:
        sentence = render_claim_sentence(claim, proj.cards)
        claim_texts.append(
            {
                "object_id": claim.claim_id,
                "text": sentence,
                "card_ids": [str(card_id) for card_id in claim.cards],
            }
        )
        citation_ids = [
            card_id
            for card_id in claim.cards
            if card_id in proj.cards and proj.cards[card_id].kind == "citation"
        ]
        if citation_ids:
            citation_texts.append(
                {
                    "object_id": claim.claim_id,
                    "text": sentence,
                    "card_ids": citation_ids,
                }
            )

    objects = table_objects + [
        {
            "object_type": "claim",
            "object_id": claim.claim_id,
            "card_ids": [str(card_id) for card_id in claim.cards],
        }
        for claim in active_claims
    ]
    for index, item in enumerate(figure_items or []):
        card_id = str(item.get("card_id") or "")
        objects.append(
            {
                "object_type": "figure",
                "object_id": str(item.get("object_id") or f"figure:{index}"),
                "card_ids": [card_id] if card_id else [],
            }
        )
    preflight = preflight_delivery(
        proj,
        objects=objects,
        tables=[table_model] if table_model is not None else None,
        claim_texts=claim_texts,
        citation_texts=citation_texts or None,
        library=library,
        figure_items=figure_items,
    )
    if not preflight.ok or preflight.manifest is None:
        codes = ", ".join(issue.code for issue in preflight.issues[:8])
        if any(issue.code in {"card_missing", "object_card_missing"} for issue in preflight.issues):
            raise ValueError(f"交付 preflight 失败：缺 EvidenceCard ({codes})")
        raise ValueError(f"交付 preflight 失败：{codes or 'evidence_not_ready'}")

    doc = Document()
    doc.add_heading(title, level=0)
    if method:
        doc.add_paragraph(method)
    if table_model is not None:
        add_table_to_doc(doc, table_model)
    for claim in active_claims:
        doc.add_paragraph(render_claim_sentence(claim, proj.cards) + f"  [{claim.claim_id}]")
    citation_ids = sorted(
        {
            card_id
            for claim in active_claims
            for card_id in claim.cards
            if card_id in proj.cards and proj.cards[card_id].kind == "citation"
        }
    )
    if citation_ids:
        doc.add_paragraph("引用文献（citable 块）")
        for card_id in citation_ids:
            loc = proj.cards[card_id].locator or {}
            doc.add_paragraph(f"- {loc.get('doc_id', '?')} p{loc.get('page', '?')}  {card_id}")
    if figure_items:
        from .figure import add_figures_to_doc

        add_figures_to_doc(doc, figure_items, proj.cards, require_card_validation=True)
    if limits:
        doc.add_paragraph("局限：" + limits)

    import io

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), preflight.manifest


def draft_from_ledger(
    proj,
    *,
    title: str = "实证研究初稿",
    method: str = "",
    limits: str = "",
    library=None,
    figure_items: list[dict] | None = None,
) -> bytes:
    """Compatibility wrapper returning only DOCX bytes."""

    data, _manifest = build_draft_package(
        proj,
        title=title,
        method=method,
        limits=limits,
        library=library,
        figure_items=figure_items,
    )
    return data


def draft_docx(results: dict[str, dict], *, title: str = "实证研究初稿",
               table1_path: str | None = None) -> bytes:
    doc = Document()
    doc.add_heading(title, level=0)
    if table1_path:
        t1 = table1_from_long(table1_path)
        if t1:
            add_table_to_doc(doc, t1)
    add_table_to_doc(doc, regression_table_from_results(results))
    rb = check_robustness(results)
    doc.add_paragraph("稳健性小结：" + ("各口径符号与显著性一致。" if rb["stable"] else "存在口径不一致，需逐行核对（见下表）。"))
    import io

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
