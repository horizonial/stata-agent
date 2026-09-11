"""figure 证据（audit C1 / DD-01 §2.6 figure 卡，A 组第三件）。

Stata 图落 artifact → 以 validator 身份签成 figure 卡（locator: run_id/path/kind）；
Word 渲染：插图 + caption（引卡号）。图不落地卡就不许进稿结论。
"""

from __future__ import annotations

import hashlib
from pathlib import Path
from collections.abc import Mapping

from ..domain.models import EvidenceCard
from ..events.schema import EVENT_CARD_SIGNED, ACTOR_VALIDATOR, Event
from ..storage.sqlite_store import SQLiteStore


def _artifact_digest(path: Path) -> tuple[str, int]:
    digest = hashlib.sha256()
    size = 0
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            size += len(block)
            digest.update(block)
    return digest.hexdigest(), size


def _is_embeddable_image(path: Path) -> bool:
    """Use python-docx's decoder as the single supported image sniff."""

    try:
        from docx.image.image import Image

        Image.from_file(str(path))
        return True
    except Exception:  # noqa: BLE001 - an unknown format is not deliverable
        return False


def sign_figure_card(
    store: SQLiteStore,
    run_id: str,
    figure_path: str | Path,
    *,
    idea: str = "i1",
    kind: str = "event-study",
    caption: str = "",
) -> str:
    """把一张已导出的图签成 figure 卡（须 run 成功、文件存在）。"""
    proj = store.project(idea)
    rec = proj.runs.get(run_id)
    if rec is None or rec.status != "succeeded":
        raise ValueError(f"run {run_id!r} 不存在或未成功，图不能作为证据")
    from ..tools.evidence_signer import validate_run_provenance

    provenance_kind = validate_run_provenance(rec.provenance or {})
    path = Path(figure_path)
    if not path.is_file() or path.is_symlink():
        raise ValueError("图文件不存在或不是 regular file")
    path = path.resolve()
    if not _is_embeddable_image(path):
        raise ValueError("图文件不是可嵌入的图片")
    artifact_sha256, byte_size = _artifact_digest(path)
    card_id = "card-fig-" + hashlib.sha256(
        f"{run_id}|{kind}|{artifact_sha256}".encode("utf-8")
    ).hexdigest()[:24]
    existing = proj.cards.get(card_id)
    if existing is not None:
        locator = existing.locator if isinstance(existing.locator, dict) else {}
        if (
            existing.kind == "figure"
            and locator.get("artifact_sha256") == artifact_sha256
            and locator.get("byte_size") == byte_size
        ):
            return card_id
        raise ValueError("figure card identity collision")
    card = EvidenceCard(
        card_id=card_id,
        kind="figure",
        locator={"run_id": run_id, "path": str(path), "kind": kind,
                 "provenance_kind": provenance_kind,
                 "artifact_sha256": artifact_sha256, "byte_size": byte_size},
        value={"caption": caption},
        signed_by=ACTOR_VALIDATOR,
    )
    store.append(Event(idea_id=idea, event_type=EVENT_CARD_SIGNED, actor=ACTOR_VALIDATOR,
                       source=ACTOR_VALIDATOR, payload={"card": card.model_dump()}))
    return card_id


def validate_figure_item(
    item: Mapping[str, object],
    cards: Mapping[str, EvidenceCard],
    *,
    require_card: bool = True,
) -> str | None:
    """Validate one figure item without exposing filesystem details."""

    card_id = str(item.get("card_id") or "")
    if not card_id:
        return "missing_card"
    card = cards.get(card_id)
    if card is None or card.kind != "figure" or not isinstance(card.locator, dict):
        return "card_missing_or_wrong_kind"
    path = Path(str(item.get("path") or ""))
    if not path.is_file() or path.is_symlink():
        return "artifact_missing_or_not_regular"
    path = path.resolve()
    locator = card.locator
    if locator.get("run_id") != item.get("run_id", locator.get("run_id")):
        return "run_mismatch"
    if "caption" in item:
        card_value = card.value if isinstance(card.value, dict) else {}
        if str(item.get("caption") or "") != str(card_value.get("caption") or ""):
            return "caption_mismatch"
    try:
        recorded_path = Path(str(locator.get("path") or "")).resolve()
    except (OSError, RuntimeError, ValueError):
        return "path_mismatch"
    if recorded_path != path:
        return "path_mismatch"
    try:
        digest, size = _artifact_digest(path)
    except (OSError, ValueError):
        return "artifact_unreadable"
    if locator.get("artifact_sha256") != digest or locator.get("byte_size") != size:
        return "artifact_digest_mismatch"
    if not _is_embeddable_image(path):
        return "artifact_not_image"
    return None


def add_figures_to_doc(
    doc,
    items: list[dict],
    cards: Mapping[str, EvidenceCard] | None = None,
    *,
    require_card_validation: bool = False,
) -> None:
    """Append validated figure items to an existing python-docx document."""

    from docx.shared import Inches

    for n, item in enumerate(items, start=1):
        card_id = item.get("card_id")
        if not card_id:
            raise ValueError("图缺 figure EvidenceCard provenance，不能出稿")
        path = Path(str(item.get("path") or ""))
        if not path.is_file() or path.is_symlink():
            raise ValueError("图文件不存在或不是 regular file")
        path = path.resolve()
        if not _is_embeddable_image(path):
            raise ValueError("图文件不是可嵌入的图片")
        if cards is not None or require_card_validation:
            error = validate_figure_item(item, cards or {}, require_card=True)
            if error:
                raise ValueError(f"figure provenance invalid: {error}")
        doc.add_picture(str(path), width=Inches(5.2))
        cap = item.get("caption") or ""
        if cards is not None:
            card = cards.get(str(card_id))
            if card is not None and isinstance(card.value, dict):
                cap = card.value.get("caption") or cap
        doc.add_paragraph(f"图 {n}：{cap}" + f"　[{card_id}]")


def figure_docx(
    title: str,
    items: list[dict],
    *,
    cards: Mapping[str, EvidenceCard] | None = None,
    require_card_validation: bool = False,
) -> bytes:
    """items: [{path, caption, card_id, claim_id?}] → 插图 Word。

    返回 .docx 字节；每图下 caption 带 [图 卡id] 溯源注。
    """
    from docx import Document
    doc = Document()
    doc.add_heading(title, level=0)
    add_figures_to_doc(
        doc,
        items,
        cards,
        require_card_validation=require_card_validation,
    )
    import io

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
