"""Wave-1：ExperimentFamily 登记/选主 + 稳健性一致 + 多表初稿 .docx。"""

from __future__ import annotations

from docx import Document

from stata_agent.domain.family import family_view, register_run, select_main
from stata_agent.storage.sqlite_store import SQLiteStore
from stata_agent.tools.fake_executor import FakeExecutor, default_test_contract
from stata_agent.tools.evidence_signer import sign_run_numeric_cards
from stata_agent.tools.robustness import check_robustness
from stata_agent.writer.draft_multi import draft_docx, draft_from_ledger, regression_table_from_results


def _results():
    return {
        "main": {"ok": True, "label": "主回归",
                 "machine": {"coef": 2.81, "se": 0.70, "N": 788, "r2": 0.008}},
        "robust_no_mgr": {"ok": True, "label": "不含管理者",
                          "machine": {"coef": 2.96, "se": 0.72, "N": 795, "r2": 0.008}},
        "robust_cluster_chain": {"ok": True, "label": "按连锁聚类",
                                 "machine": {"coef": 2.81, "se": 1.40, "N": 788, "r2": 0.008}},
    }


def _run_family(store):
    """用 FakeExecutor 跑 3 个 spec 并登记进同一 family、选主结果（等价旧 engine 的语义）。"""
    ex = FakeExecutor(store)
    for vid in ("main", "robust_no_mgr", "robust_cluster_chain"):
        out = ex.execute(
            f"sysuse auto, clear\nreg price mpg\n* variant {vid}",
            idea="i1",
            run_id=f"i1-{vid}",
            result_contract=default_test_contract(),
        )
        sign_run_numeric_cards(store, out["run_id"], idea="i1")
        register_run(store, "i1", "family-i1", out["run_id"], variant=vid)
    select_main(store, "i1", "family-i1", "i1-main", reason="主回归")
    return store


def test_family_registers_all_and_selects_main(tmp_path):
    store = _run_family(SQLiteStore(str(tmp_path / "l.db"), writer_id="a"))
    view = family_view(store, "i1", "family-i1")
    assert len(view["members"]) == 3
    assert view["main"] is not None and view["main"]["run_id"] == "i1-main"
    store.close()


def test_robustness_stable_and_flips():
    rb = check_robustness(_results())
    assert rb["stable"] is True
    # 造一个反号变体 → 不稳
    flipped = dict(_results())
    flipped["robust_no_mgr"] = dict(flipped["robust_no_mgr"])
    flipped["robust_no_mgr"]["machine"] = dict(flipped["robust_no_mgr"]["machine"], coef=-1.0)
    assert check_robustness(flipped)["stable"] is False


def test_draft_from_ledger(tmp_path):
    store = _run_family(SQLiteStore(str(tmp_path / "l.db"), writer_id="a"))
    proj = store.project("i1")
    data = draft_from_ledger(proj, method="方法段", limits="局限段")
    p = tmp_path / "ledger.docx"
    p.write_bytes(data)
    doc = Document(str(p))
    paras = [par.text for par in doc.paragraphs]
    assert any("方法段" in t for t in paras) and any("局限段" in t for t in paras)
    allcell = " | ".join(c.text for tb in doc.tables for row in tb.rows for c in row.cells)
    assert "样本量" in allcell
    store.close()


def test_draft_docx_tables(tmp_path):
    model = regression_table_from_results(_results())
    assert [row.label for row in model.rows] == ["系数", "标准误", "样本量", "R²"]
    assert model.columns == ["主回归", "不含管理者", "按连锁聚类"]
    # 主系数带星
    assert model.rows[0].cells[0].text.startswith("2.810")

    data = draft_docx(_results())
    p = tmp_path / "draft.docx"
    p.write_bytes(data)
    doc = Document(str(p))
    paras = [par.text for par in doc.paragraphs]
    assert any("稳健性小结" in t for t in paras)
    allcell = " | ".join(c.text for tb in doc.tables for row in tb.rows for c in row.cells)
    assert "主回归" in allcell and "标准误" in allcell
